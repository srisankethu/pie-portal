"""One catalogue out of several files, and the prices that must not survive it.

``test_company_catalogues.py`` pins the per-company shape: one company, its own
export, its own pack. These pin the shape *inside* one company — an item master,
a manufacturer's range extension and a price list, merged into one catalogue —
and the rules that merging brings with it.

Three of them are the ones worth reading:

* **a price never reaches the catalogue**, asserted on the stored corpus and on
  the decoded output rather than on the ingest report that claims it. The report
  is the thing being tested; believing it would be circular.
* **a colliding part number resolves to the newest file, and is counted.**
  pie-parser's ``AuthoritativeIndex`` treats a duplicate identifier inside one
  namespace as a collision that never resolves, so an unmerged duplicate does
  not give a wrong answer — it silently stops that part number resolving at all,
  which no record count would reveal.
* **removing a source does not rebuild anything.** The catalogue on disk keeps
  resolving and goes out of date; a build that happened as a side effect of
  tidying a file list would replace what a whole company resolves against.
"""
from __future__ import annotations

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import sessionmaker

import dbsupport
from app import catalog
from app.config import settings
from app.db import get_session
from app.domain import models
from app.routers import data_status, platform_auth
from app.seed import SEED_PASSWORD, ensure_org_and_users

requires_pie = pytest.mark.requires_pie

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: A price list as one actually arrives: a title row before the headings, the
#: part number under a name the pack has never heard of, and money beside it.
PRICE = 4210.50


@pytest.fixture()
def client():
    engine = dbsupport.fresh_engine()
    Maker = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False,
                         future=True)
    s = Maker()
    ensure_org_and_users(s)
    s.add(models.ZohoConnection(connection_id="cx_sls", organization_id="org_pie",
                                label="SLS Engineers", zoho_organization_id="z1"))
    s.commit()
    s.close()

    app = FastAPI()
    app.include_router(platform_auth.router)
    app.include_router(data_status.router)

    def _override():
        sess = Maker()
        try:
            yield sess
            sess.commit()
        finally:
            sess.close()

    app.dependency_overrides[get_session] = _override
    tc = TestClient(app)
    tc.Maker = Maker
    return tc


def _hdr(c, email="s.menon@pie.example"):
    r = c.post("/api/v1/auth/login", json={"email": email, "password": SEED_PASSWORD})
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _upload(c, hdr, content, filename, source_key=None, ctype="text/csv"):
    url = (f"/api/v1/data/catalog/companies/cx_sls/corpus"
           f"?filename={filename}")
    if source_key is not None:
        url += f"&source_key={source_key}"
    return c.post(url, content=content, headers={**hdr, "Content-Type": ctype})


def _price_list_xlsx(rows, headers=("Part No", "Item Description", "Grade",
                                    "New ZCNC Price", "Stock Qty")) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["ACME price list — revision 3", None, None])
    sheet.append([None, None, None])
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _sources(body):
    return {s["source_key"]: s for s in body["sources"]}


# ── the nomenclature-only invariant, checked on the artefacts ────────────────

@requires_pie
def test_no_price_survives_from_an_uploaded_price_list(client, tmp_path,
                                                       monkeypatch):
    """A price list decodes into nomenclature and nothing else.

    Asserted on the corpus the build read and on the decoded JSONL, not on the
    ingest report — the report is what claims the columns were dropped, so
    trusting it here would be testing the claim against itself.

    The corpus is checked *as the parse saw it*: the uploaded bytes are stored
    whole, deliberately, so the catalogue can be rebuilt after a redeploy. What
    must never carry a price is the corpus the pack reads and the catalogue it
    writes, and both are reconstructed here the way ``build_for_company`` does.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    raw = _price_list_xlsx([
        [1234567, "SC DRILL 8.00MM 5XD COOLANT", "KC7315", PRICE, 12],
        [7654321, "INSERT ANSI/ISO TURNING CNMG 120408", "KCP25B", 812.00, 40],
    ])
    body = _upload(client, hdr, raw, "prices.xlsx", "prices.xlsx", XLSX).json()
    assert body["sources"][0]["ingest"]["commercial_columns_dropped"] == [
        "New ZCNC Price", "Stock Qty"]

    built = client.post("/api/v1/data/catalog/companies/cx_sls/build",
                        headers=hdr).json()
    assert built["records"] == 2

    s = client.Maker()
    catalog.combined_corpus(catalog.current_corpora(s, "org_pie", "cx_sls"),
                            settings.PIE_PACK, tmp_path / "merged.csv")
    s.close()
    merged = (tmp_path / "merged.csv").read_bytes()
    # The header is the pack's, and the money columns are not in it at all —
    # absent because they were never written, not filtered afterwards.
    assert merged.decode().splitlines()[0] == "MM#,Material Description,Grade"
    assert b"4210" not in merged
    assert b"New ZCNC Price" not in merged and b"Stock Qty" not in merged

    decoded = catalog.company_catalog_path("cx_sls").read_text(encoding="utf-8")
    # `"4210.5"` rather than `4210`: the loose form matches inside unrelated
    # part numbers (2984210) and would pass against a catalogue that really did
    # carry the price. The exact value, and the column name, are the claim.
    assert "4210.5" not in decoded
    assert "New ZCNC Price" not in decoded
    assert "SC DRILL 8.00MM 5XD COOLANT" in decoded


@requires_pie
def test_an_excel_export_decodes_without_being_converted_first(client, tmp_path,
                                                               monkeypatch):
    """The part number survives Excel's own typing.

    A numeric part number comes back from openpyxl as a float, and ``str()`` on
    it gives ``1234567.0`` — a record id that matches nothing a customer will
    ever type, in a catalogue that otherwise looks perfectly built. This is the
    reason the reader converts cells itself rather than formatting them at the
    edge.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    raw = _price_list_xlsx([[1234567, "SC DRILL 8.00MM 5XD COOLANT", "KC7315",
                             PRICE, 1]])
    _upload(client, hdr, raw, "prices.xlsx", "prices.xlsx", XLSX)
    client.post("/api/v1/data/catalog/companies/cx_sls/build", headers=hdr)

    decoded = catalog.company_catalog_path("cx_sls").read_text(encoding="utf-8")
    # `dump_stable_json` writes compact separators, so the key and value sit
    # together — asserted as one string rather than on the number alone, which
    # would also match a dimension token that happens to contain it.
    assert '"record_id":"1234567"' in decoded
    assert "1234567.0" not in decoded


@requires_pie
def test_a_manufacturers_price_list_shape_decodes(client, tmp_path, monkeypatch):
    """The file as it actually arrives, preamble and totals row included.

    Every element here broke something while this was being written. The header
    is not row one — there is a title, a blank line and an ``Effective`` date
    above it, and taking the first row with two filled cells as the headings
    picked the date row and made every real column unmappable. The totals row
    has no part number and must be skipped rather than decoded. And the price
    column must be named as ignored, because that is the only form in which
    "nomenclature only" is a claim the person who uploaded the file can check.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)

    book = Workbook()
    sheet = book.active
    sheet.append(["KENNAMETAL INDIA — PRICE LIST", None, None, None])
    sheet.append([None, None, None, None])
    sheet.append(["Effective", "01-04-2026", None, None])
    sheet.append(["Part No", "Description", "Grade", "List Price (INR)"])
    sheet.append([1234567, "SC DRILL 8.00MM 5XD COOLANT", "KC7315", PRICE])
    sheet.append([None, "TOTAL", None, PRICE])
    buf = io.BytesIO()
    book.save(buf)

    body = _upload(client, hdr, buf.getvalue(), "kmt.xlsx", "kmt", XLSX).json()
    ingest = _sources(body)["kmt"]["ingest"]
    assert ingest["mapped"]["record_id"] == "Part No"
    assert ingest["mapped"]["description"] == "Description"
    assert ingest["rows_kept"] == 1
    assert ingest["rows_skipped_blank_key"] == 1          # the totals row
    assert ingest["commercial_columns_dropped"] == ["List Price (INR)"]

    built = client.post("/api/v1/data/catalog/companies/cx_sls/build",
                        headers=hdr).json()
    assert built["records"] == 1
    decoded = catalog.company_catalog_path("cx_sls").read_text(encoding="utf-8")
    assert "4210.5" not in decoded
    assert '"record_id":"1234567"' in decoded


def test_only_the_first_sheet_of_a_workbook_is_read(client):
    """A "Discontinued" tab must not be merged into the catalogue.

    Which sheet of a multi-sheet workbook is the item master is a question this
    cannot answer, and answering it wrongly puts withdrawn products into what a
    quote resolves against. Reading one sheet is the conservative half of that
    pair, and it is logged rather than guessed at.
    """
    hdr = _hdr(client)
    book = Workbook()
    current = book.active
    current.title = "Current"
    current.append(["MM#", "Material Description", "Grade"])
    current.append(["A", "SC DRILL 6.00MM 3XD", "KC7315"])
    old_tab = book.create_sheet("Discontinued")
    old_tab.append(["MM#", "Material Description", "Grade"])
    old_tab.append(["Z", "WITHDRAWN TOOL", "KC725M"])
    buf = io.BytesIO()
    book.save(buf)

    body = _upload(client, hdr, buf.getvalue(), "two.xlsx", "two", XLSX).json()
    assert _sources(body)["two"]["ingest"]["rows_kept"] == 1


def test_a_large_export_is_read_without_being_held_in_memory(client):
    """The rows are a stream, and this is what says so.

    Materialising a 33 MB CSV as lists of cells peaked at 394 MB resident — for
    one file, on an upload an owner can repeat, on a container sized in
    hundreds of megabytes. Asserted structurally rather than by measuring RSS,
    which is neither stable across machines nor meaningful under a shared
    interpreter: ``Table`` must not hold a materialised row list, and iterating
    twice must give the same rows both times, which a consumed iterator would
    not.
    """
    from app.ingestion import item_master

    raw = b"MM#,Material Description,Grade\n" + b"".join(
        f"A{i},SC DRILL {i}.00MM 5XD,KC7315\n".encode() for i in range(500))
    table = item_master.read_table(raw, "big.csv")

    # Nothing on the table is a list of rows — only the header and the bytes.
    held = {name: getattr(table, name) for name in table.__slots__}
    assert not any(isinstance(v, list) and v and isinstance(v[0], list)
                   for v in held.values())

    first = list(table.rows())
    second = list(table.rows())
    assert len(first) == 500
    assert first == second        # re-readable, not a consumed iterator

    report = item_master.ingest_report(table, item_master.suggest_mapping(table))
    rows = item_master.emit_rows(table, {"record_id": "MM#",
                                         "description": "Material Description",
                                         "grade": "Grade"}, report)
    # A generator: the counts are only true once it has been exhausted, which
    # is the contract `describe` and `combined_corpus` both rely on.
    assert "rows_kept" not in report
    assert len(list(rows)) == 500
    assert report["rows_kept"] == 500


# ── several files, one catalogue ─────────────────────────────────────────────

def test_a_named_source_replaces_only_itself(client):
    """Two files stay two files; re-uploading one leaves the other alone.

    Without a key, the second upload superseded the first — which is right for
    "replace the export" and wrong for "add the price list", and the difference
    is a company silently resolving against one of its two masters.
    """
    hdr = _hdr(client)
    header = b"MM#,Material Description,Grade\n"
    _upload(client, hdr, header + b"A,FIRST TOOL,KC725M\n", "master.csv", "master")
    body = _upload(client, hdr, header + b"B,SECOND TOOL,KC725M\n",
                   "prices.csv", "prices").json()
    assert sorted(_sources(body)) == ["master", "prices"]

    again = _upload(client, hdr, header + b"B,SECOND TOOL REVISED,KC725M\n",
                    "prices.csv", "prices").json()
    assert sorted(_sources(again)) == ["master", "prices"]
    assert _sources(again)["prices"]["corpus_id"] != \
        _sources(body)["prices"]["corpus_id"]

    # And the unkeyed form still means what it always meant: replace the lot.
    whole = _upload(client, hdr, header + b"C,ONLY TOOL,KC725M\n",
                    "fresh.csv").json()
    assert list(_sources(whole)) == ["fresh.csv"]


def test_removing_a_source_leaves_the_built_catalogue_alone(client, tmp_path,
                                                            monkeypatch):
    """Out of date, not rebuilt, and not gone.

    A rebuild here would replace what a whole company resolves against as a
    side effect of tidying a list of files. The catalogue keeps its records and
    its stamp; ``stale`` is what says it no longer matches what is on record.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    header = b"MM#,Material Description,Grade\n"
    _upload(client, hdr, header + b"A,FIRST TOOL,KC725M\n", "master.csv", "master")
    _upload(client, hdr, header + b"B,SECOND TOOL,KC725M\n", "extra.csv", "extra")

    after = client.delete(
        "/api/v1/data/catalog/companies/cx_sls/sources/extra", headers=hdr).json()
    assert list(_sources(after)) == ["master"]
    # Nothing was built in this test, so there is nothing to be stale; the point
    # is that the removal did not build one either.
    assert after["exists"] is False
    assert after["built_at"] is None

    gone = client.delete(
        "/api/v1/data/catalog/companies/cx_sls/sources/extra", headers=hdr)
    assert gone.status_code == 404


@requires_pie
def test_a_removed_source_marks_the_catalogue_out_of_date(client, tmp_path,
                                                          monkeypatch):
    """The staleness a single ``corpus_id`` comparison could not see.

    Removing the *older* of two sources leaves the newest corpus id untouched,
    so the old check called the catalogue current while a rebuild would have
    produced a different one. Comparing digests over the set is what fixes it,
    and this is the case that distinguishes the two.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    header = b"MM#,Material Description,Grade\n"
    _upload(client, hdr, header + b"A,SC DRILL 6.00MM 3XD,KC7315\n",
            "master.csv", "master")
    _upload(client, hdr, header + b"B,SC DRILL 8.00MM 5XD,KC7315\n",
            "extra.csv", "extra")
    built = client.post("/api/v1/data/catalog/companies/cx_sls/build",
                        headers=hdr).json()
    assert built["records"] == 2
    assert built["stale"] is False
    assert len(built["built_from"]) == 2

    after = client.delete(
        "/api/v1/data/catalog/companies/cx_sls/sources/master",
        headers=hdr).json()
    assert after["stale"] is True
    assert after["exists"] is True            # still a real catalogue
    assert after["records"] == 2              # and still the one that was built

    rebuilt = client.post("/api/v1/data/catalog/companies/cx_sls/build",
                          headers=hdr).json()
    assert rebuilt["stale"] is False
    assert rebuilt["records"] == 1


@requires_pie
def test_the_newest_file_wins_a_collision_and_it_is_counted(client, tmp_path,
                                                            monkeypatch):
    """The rule that keeps a duplicated part number resolving at all.

    Two files carrying ``A`` must not both emit it: pie-parser indexes
    identifiers per namespace and a duplicate inside one namespace resolves for
    neither. So the newest file's row is the one emitted — a later file is a
    later statement about the same product — and the overlap is counted rather
    than absorbed, because two exports disagreeing about one product is
    something a person has to be told.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    header = b"MM#,Material Description,Grade\n"
    _upload(client, hdr, header + b"A,SC DRILL 6.00MM 3XD,KC7315\n",
            "master.csv", "master")
    _upload(client, hdr, header + b"A,SC DRILL 8.00MM 5XD,KC7315\n",
            "revised.csv", "revised")

    built = client.post("/api/v1/data/catalog/companies/cx_sls/build",
                        headers=hdr).json()
    assert built["records"] == 1
    assert built["ingest"]["collisions"] == 1
    assert built["ingest"]["collision_examples"] == ["A"]

    decoded = catalog.company_catalog_path("cx_sls").read_text(encoding="utf-8")
    assert "8.00MM" in decoded            # the newer statement
    assert "6.00MM" not in decoded


# ── the mapping: a file keeps its own column names ───────────────────────────

@requires_pie
def test_a_file_with_its_own_column_names_builds_after_being_mapped(
        client, tmp_path, monkeypatch):
    """The case that used to be unbuildable.

    A master with two code columns is exactly where the guess goes wrong, and
    where the old behaviour — refuse anything not phrased like the pack —
    offered no way forward at all. The mapping is corrected against the file's
    own headers and the build then reads the intended column.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    raw = ("Item Code,Legacy Ref,Particulars,Grade\n"
           "NEW-1,OLD-1,SC DRILL 8.00MM 5XD COOLANT,KC7315\n").encode()
    body = _upload(client, hdr, raw, "odd.csv", "odd").json()
    guessed = _sources(body)["odd"]["mapping"]
    assert guessed["record_id"] == "Item Code"
    assert guessed["description"] == "Particulars"

    fixed = client.put(
        "/api/v1/data/catalog/companies/cx_sls/sources/odd/mapping",
        json={"record_id": "Legacy Ref", "description": "Particulars",
              "grade": "Grade"}, headers=hdr).json()
    assert _sources(fixed)["odd"]["mapping"]["record_id"] == "Legacy Ref"
    # The re-read reports what the new mapping leaves out, so the change is
    # visible rather than merely stored.
    assert "Item Code" in _sources(fixed)["odd"]["ingest"]["dropped_columns"]

    client.post("/api/v1/data/catalog/companies/cx_sls/build", headers=hdr)
    decoded = catalog.company_catalog_path("cx_sls").read_text(encoding="utf-8")
    assert "OLD-1" in decoded and "NEW-1" not in decoded


def test_a_build_names_the_file_it_could_not_read(client):
    """"The build failed" over six files is not something a person can act on.

    Reached by storing a mapping and then replacing the file with one that has
    different headers under the same key — which is the realistic way a source
    stops being readable, and the case where naming the file is the whole of the
    answer. Reported as the file's own problem rather than as a parser failure,
    because nothing reached the parser.
    """
    hdr = _hdr(client)
    client.put("/api/v1/data/catalog/companies/cx_sls/pack",
               json={"pack_id": "zcnc"}, headers=hdr)
    _upload(client, hdr, b"MM#,Material Description,Grade\nA,A TOOL,KC725M\n",
            "master.csv", "master")
    _upload(client, hdr, b"Item Code,Particulars\nB,B TOOL\n",
            "extra.csv", "extra")
    client.put("/api/v1/data/catalog/companies/cx_sls/sources/extra/mapping",
               json={"record_id": "Item Code", "description": "Particulars"},
               headers=hdr)

    # The same source key, a file that no longer has the mapped columns. The
    # upload itself is fine — its own headers are readable — but the stored
    # mapping now names columns this file lacks.
    s = client.Maker()
    row = next(r for r in s.query(models.CompanyCorpus).all()
               if r.source_key == "extra" and r.superseded_at is None)
    row.content = b"Other,Columns\nC,C TOOL\n"
    s.commit()
    s.close()

    r = client.post("/api/v1/data/catalog/companies/cx_sls/build", headers=hdr)
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert detail.startswith("extra.csv:")
    assert "Item Code" in detail


def test_a_mapping_naming_a_column_the_file_lacks_is_refused(client):
    """Refused rather than stored: a mapping that cannot build would leave the
    company unable to build with no statement of why."""
    hdr = _hdr(client)
    raw = b"MM#,Material Description,Grade\nA,FIRST TOOL,KC725M\n"
    _upload(client, hdr, raw, "master.csv", "master")
    r = client.put(
        "/api/v1/data/catalog/companies/cx_sls/sources/master/mapping",
        json={"record_id": "Nope", "description": "Material Description"},
        headers=hdr)
    assert r.status_code == 422
    assert "Nope" in r.json()["detail"]
    assert "MM#" in r.json()["detail"]        # what the file does have


def test_a_workbook_named_as_a_csv_is_told_what_it_is(client):
    """The mislabelled-file case, named rather than reported as a broken CSV.

    Read as text, a workbook's zip container produces a parse failure thousands
    of characters in that says nothing about the actual problem.
    """
    hdr = _hdr(client)
    raw = _price_list_xlsx([[1, "A TOOL", "KC725M", 1.0, 1]])
    r = _upload(client, hdr, raw, "prices.csv", "prices")
    assert r.status_code == 422
    assert "Excel" in r.json()["detail"]


def test_the_number_of_sources_is_capped_by_name(client, monkeypatch):
    """A ceiling nobody should reach, refused with what to do about it: every
    source is read on every rebuild, so this is a real cost rather than a
    formality."""
    monkeypatch.setattr(catalog, "MAX_SOURCES", 2)
    hdr = _hdr(client)
    header = b"MM#,Material Description,Grade\n"
    for i in range(2):
        assert _upload(client, hdr, header + f"A{i},A TOOL,KC725M\n".encode(),
                       f"f{i}.csv", f"f{i}").status_code == 200
    third = _upload(client, hdr, header + b"A3,A TOOL,KC725M\n", "f3.csv", "f3")
    assert third.status_code == 409
    assert "limit of 2" in third.json()["detail"]

    # Replacing one of the two is still allowed at the ceiling — the limit is on
    # how many files are kept, not on how often they are updated.
    assert _upload(client, hdr, header + b"A0,A BETTER TOOL,KC725M\n",
                   "f0.csv", "f0").status_code == 200


# ── the pack trial ──────────────────────────────────────────────────────────

@requires_pie
def test_the_pack_trial_reports_the_parsers_counts_and_writes_nothing(
        client, tmp_path, monkeypatch):
    """Choosing a pack becomes a decision with evidence behind it.

    Two things are asserted: the counts come back per pack, and the trial leaves
    no catalogue on disk. The second matters more than it looks — a trial that
    wrote its output where a build writes it would leave a company resolving
    against a pack it never chose, stamped as provenanced.
    """
    monkeypatch.setattr(settings, "PIE_CATALOG", tmp_path / "products.jsonl")
    hdr = _hdr(client)
    header = b"MM#,Material Description,Grade\n"
    _upload(client, hdr, header + b"A,SC DRILL 8.00MM 5XD COOLANT,KC7315\n",
            "master.csv", "master")

    fit = client.get("/api/v1/data/catalog/companies/cx_sls/pack-fit",
                     headers=hdr).json()
    assert fit["available"] is True
    assert [p["pack_id"] for p in fit["packs"]] == ["zcnc"]
    assert fit["packs"][0]["rows_read"] == 1
    assert fit["packs"][0]["classified"] == 1
    assert not catalog.company_catalog_path("cx_sls").exists()


def test_the_pack_trial_says_why_it_has_nothing_to_try(client):
    """``available: False`` with a reason, never an empty list.

    An empty result would read as "no pack fits your file", which is a claim
    about the file. The two real causes — no file uploaded, no pack shipped —
    are different problems with different fixes.
    """
    hdr = _hdr(client)
    fit = client.get("/api/v1/data/catalog/companies/cx_sls/pack-fit",
                     headers=hdr).json()
    assert fit["available"] is False
    assert fit["packs"] == []
    assert "no item-master export" in fit["reason"]


def test_the_source_actions_are_owner_only(client):
    """Same line as the build: reading which catalogue answered is open, and
    changing what a company resolves against is not."""
    hdr = _hdr(client)
    raw = b"MM#,Material Description,Grade\nA,FIRST TOOL,KC725M\n"
    _upload(client, hdr, raw, "master.csv", "master")

    for email in ("r.nair@pie.example", "m.rao@pie.example"):
        other = _hdr(client, email)
        assert client.delete(
            "/api/v1/data/catalog/companies/cx_sls/sources/master",
            headers=other).status_code == 403
        assert client.put(
            "/api/v1/data/catalog/companies/cx_sls/sources/master/mapping",
            json={"record_id": "MM#", "description": "Material Description"},
            headers=other).status_code == 403
        assert client.get("/api/v1/data/catalog/companies/cx_sls/pack-fit",
                          headers=other).status_code == 403
        # Reading stays open.
        assert client.get("/api/v1/data/catalog/companies",
                          headers=other).status_code == 200


def test_a_source_of_another_organization_reads_as_absent(client):
    """A key from another tenant is a 404, not a 403 that confirms it exists."""
    hdr = _hdr(client)
    s = client.Maker()
    s.add(models.Organization(organization_id="org_else", name="Someone Else"))
    s.add(models.ZohoConnection(connection_id="cx_other",
                                organization_id="org_else",
                                label="Someone Else", zoho_organization_id="z9"))
    s.add(models.CompanyCorpus(organization_id="org_else",
                               connection_id="cx_other",
                               source_key="theirs", filename="theirs.csv",
                               size_bytes=1, sha256="x", content=b"x"))
    s.commit()
    s.close()

    assert client.delete(
        "/api/v1/data/catalog/companies/cx_other/sources/theirs",
        headers=hdr).status_code == 404
